import os
from openpyxl import Workbook, load_workbook

FILE = "accounts.xlsx"


def load_book():
    """Open the Excel file if it exists, otherwise create a new one."""
    if os.path.exists(FILE):
        return load_workbook(FILE)
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Phone No", "Balance"])
    wb.save(FILE)
    return wb


def save_book(wb):
    """Save the workbook. Returns True if saved, False if the file is open in Excel."""
    try:
        wb.save(FILE)
        return True
    except PermissionError:
        print("The Excel file is open! Please close it and try again.")
        return False


def find_row(ws, phone):
    """Find the row of an account using the phone number. Returns None if not found."""
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=2).value) == phone:
            return r
    return None


def get_name():
    """Ask for the full name (at least first name and last name, letters only)."""
    name = input("Enter full name (first name and last name): ").strip()
    parts = name.split()
    if len(parts) < 2:
        print("Please enter your FULL name (first name and last name).")
        return None
    for part in parts:
        if not part.isalpha():
            print("Name must contain only letters.")
            return None
    return " ".join(part.capitalize() for part in parts)


def get_phone():
    """Ask for a phone number that is exactly 10 digits."""
    phone = input("Enter 10-digit phone number: ").strip()
    if not phone.isdigit() or len(phone) != 10:
        print("Phone number must be exactly 10 digits (numbers only).")
        return None
    return phone


def get_amount(message):
    """Get a valid positive amount from the user."""
    try:
        amount = float(input(message))
        if amount <= 0:
            print("Amount must be greater than 0.")
            return None
        return amount
    except ValueError:
        print("Please enter a valid number.")
        return None


def menu():
    while True:
        print("\n=== Welcome to ATM System ===")
        print("1. Create account / add new account")
        print("2. Deposit money")
        print("3. Withdraw money")
        print("4. View balance")
        print("5. Exit")
        choice = input("Choose an option (1-5): ").strip()

        if choice == "1":
            add_new()
        elif choice == "2":
            deposit()
        elif choice == "3":
            withdraw()
        elif choice == "4":
            view_balance()
        elif choice == "5":
            print("Thank you! Closing the ATM system.")
            break
        else:
            print("Invalid option, please try again.")


def add_new():
    wb = load_book()
    ws = wb.active

    name = get_name()
    if name is None:
        return

    phone = get_phone()
    if phone is None:
        return

    if find_row(ws, phone):
        print("This phone number is already registered.")
        return

    balance = get_amount("Enter initial deposit amount: ")
    if balance is None:
        return

    ws.append([name, phone, balance])
    if save_book(wb):
        print("Your account is successfully created.")


def deposit():
    wb = load_book()
    ws = wb.active

    phone = get_phone()
    if phone is None:
        return

    row = find_row(ws, phone)
    if row is None:
        print("Account not found.")
        return

    amount = get_amount("Enter money to deposit: ")
    if amount is None:
        return

    ws.cell(row=row, column=3).value += amount
    if save_book(wb):
        print("Your money is successfully deposited.")


def withdraw():
    wb = load_book()
    ws = wb.active

    phone = get_phone()
    if phone is None:
        return

    row = find_row(ws, phone)
    if row is None:
        print("Account not found.")
        return

    amount = get_amount("Enter money to withdraw: ")
    if amount is None:
        return

    current = ws.cell(row=row, column=3).value
    if amount > current:
        print("Not enough balance. Your balance is:", current)
        return

    ws.cell(row=row, column=3).value = current - amount
    if save_book(wb):
        print("Your money has been successfully withdrawn.")


def view_balance():
    wb = load_book()
    ws = wb.active

    phone = get_phone()
    if phone is None:
        return

    row = find_row(ws, phone)
    if row is None:
        print("Account not found.")
        return

    print("Name:", ws.cell(row=row, column=1).value)
    print("Balance:", ws.cell(row=row, column=3).value)


menu()
